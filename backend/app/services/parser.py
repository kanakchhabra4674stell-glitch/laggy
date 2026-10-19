import os
import re
from typing import List, Dict, Any, Optional


def parse_vtt_or_srt(content: str) -> List[Dict[str, Any]]:
    """
    Parses VTT or SRT string into a list of transcript segment dicts:
    [{start_seconds: float, end_seconds: float, text: str}]
    """
    segments = []
    # Normalize line endings
    content = content.replace("\r\n", "\n").replace("\r", "\n")

    # Regular expression for SRT/VTT timestamp: 00:01:23.456 or 00:01:23,456 or 01:23.456
    time_pattern = re.compile(
        r"(?:(\d{1,2}):)?(\d{2}):(\d{2})[.,](\d{3})\s*-->\s*(?:(\d{1,2}):)?(\d{2}):(\d{2})[.,](\d{3})"
    )

    blocks = content.strip().split("\n\n")
    for block in blocks:
        lines = [line.strip() for line in block.split("\n") if line.strip()]
        if not lines:
            continue

        match = None
        text_lines = []

        for line in lines:
            m = time_pattern.search(line)
            if m:
                match = m
            elif not line.isdigit() and "WEBVTT" not in line and not line.startswith("NOTE"):
                text_lines.append(line)

        if match and text_lines:
            # Parse start time
            h1 = int(match.group(1)) if match.group(1) else 0
            m1 = int(match.group(2))
            s1 = int(match.group(3))
            ms1 = int(match.group(4))
            start_sec = h1 * 3600 + m1 * 60 + s1 + ms1 / 1000.0

            # Parse end time
            h2 = int(match.group(5)) if match.group(5) else 0
            m2 = int(match.group(6))
            s2 = int(match.group(7))
            ms2 = int(match.group(8))
            end_sec = h2 * 3600 + m2 * 60 + s2 + ms2 / 1000.0

            clean_text = " ".join(text_lines).strip()
            # Remove any HTML tags like <v Instructor>
            clean_text = re.sub(r"<[^>]+>", "", clean_text).strip()

            if clean_text:
                segments.append({
                    "start_seconds": round(start_sec, 2),
                    "end_seconds": round(end_sec, 2),
                    "text": clean_text
                })

    return segments


def parse_pdf(file_path: str) -> List[Dict[str, Any]]:
    """
    Extracts text from PDF page by page using pymupdf.
    Returns: [{page_number: int, text: str}]
    """
    import pymupdf

    pages = []
    doc = pymupdf.open(file_path)
    try:
        for page_idx in range(len(doc)):
            page = doc[page_idx]
            text = page.get_text("text").strip()
            if text:
                pages.append({
                    "page_number": page_idx + 1,
                    "text": text
                })
    finally:
        doc.close()
    return pages


def parse_pptx(file_path: str) -> List[Dict[str, Any]]:
    """
    Extracts text from PPTX slide by slide using python-pptx.
    Returns: [{slide_number: int, text: str}]
    """
    from pptx import Presentation

    prs = Presentation(file_path)
    slides = []
    for idx, slide in enumerate(prs.slides):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_texts.append(line)
        full_slide_text = "\n".join(slide_texts).strip()
        if full_slide_text:
            slides.append({
                "slide_number": idx + 1,
                "text": full_slide_text
            })
    return slides


def parse_docx(file_path: str) -> str:
    """
    Extracts full text from DOCX file using python-docx.
    """
    import docx

    doc = docx.Document(file_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def parse_text_file(file_path: str) -> str:
    """
    Reads plain text or markdown file.
    """
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def chunk_text(
    text: str,
    chunk_size: int = 400,
    chunk_overlap: int = 80
) -> List[str]:
    """
    Splits long text into overlapping chunks by word boundaries.
    """
    words = text.split()
    if not words:
        return []

    if len(words) <= chunk_size:
        return [" ".join(words)]

    chunks = []
    start = 0
    while start < len(words):
        end = min(start + chunk_size, len(words))
        chunks.append(" ".join(words[start:end]))
        if end >= len(words):
            break
        start += chunk_size - chunk_overlap
    return chunks
